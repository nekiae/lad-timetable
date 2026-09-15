"""Данные школы в Excel: шаблон, который объясняет сам себя, и терпимая загрузка.

Скачанный файл и раньше годился как шаблон, но нигде не было сказано, что писать
в колонках, — и загрузка молча портила то, что люди пишут естественно:
  • «нет» в колонке «совместитель» читалось как «да» (непустой текст — истина);
  • методический день «Ср» или «среда» пропадал — ждали цифру;
  • «Спортзал» с большой буквы не узнавался как тип кабинета.

Здесь одно описание колонок (SPEC), и из него собирается всё: лист «Как заполнять»,
подсказки у заголовков, выпадающие списки, панель на экране и разбор значений при
загрузке. Что разбор не понял — не угадывается, а попадает в отчёт с листом,
строкой и колонкой.
"""

from __future__ import annotations

import io
import math
from dataclasses import asdict, dataclass, field

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .tables import DAY_NAMES, LESSON_KINDS, LEVELS, NONE_CHOICE, ROOM_KINDS, blank_tables


@dataclass
class Column:
    name: str
    kind: str  # text | number | bool | day | choice
    hint: str
    example: str = ""
    required: bool = False
    options: list[str] = field(default_factory=list)


@dataclass
class SheetSpec:
    sheet: str
    table: str
    about: str
    columns: list[Column]


ROOMS = list(ROOM_KINDS)
WEEKDAYS = [DAY_NAMES[d] for d in range(1, 6)]

SPEC = [
    SheetSpec("Классы", "classes", "Один класс — одна строка.", [
        Column("класс", "text", "Цифра года обучения и буква, без пробела.", "5А", True),
        Column("учеников", "number", "Сколько детей в классе.", "24"),
        Column("повышенный уровень", "bool", "Можно не заполнять: класс станет профильным сам, если в нагрузке "
               "есть предмет повышенного уровня.", "нет"),
    ]),
    SheetSpec("Кабинеты", "rooms", "Один кабинет — одна строка. Спортзал, где занимаются два класса сразу, — "
              "одна строка с «классов сразу» = 2.", [
        Column("кабинет", "text", "Номер или название, как на двери.", "21", True),
        Column("тип", "choice", "Какие уроки здесь можно вести.", "обычный", True, ROOMS),
        Column("мест", "number", "Сколько учеников помещается.", "30"),
        Column("классов сразу", "number", "Сколько классов занимаются одновременно. Спортзал обычно 2, "
               "остальные кабинеты 1.", "1"),
    ]),
    SheetSpec("Предметы", "subjects", "Один предмет — одна строка. Название пишется одинаково здесь "
              "и на листе «Нагрузка».", [
        Column("предмет", "text", "Название как в учебном плане.", "Математика", True),
        Column("кабинет", "choice", "Какой кабинет нужен предмету. Пусто — обычный.", "обычный", False, ROOMS),
        Column("только в нём", "bool", "да — урок не пойдёт в другой кабинет. Пусто — решит система: физкультура, "
               "информатика и труд только в своём, остальные могут и в обычном.", ""),
        Column("всегда парой", "bool", "да — два урока подряд в один день (обычно труд).", "нет"),
    ]),
    SheetSpec("Учителя", "teachers", "Один учитель — одна строка. ФИО пишется одинаково здесь "
              "и на листе «Нагрузка».", [
        Column("ФИО", "text", "Фамилия и инициалы — так, как будет в расписании.", "Иванова И. И.", True),
        Column("методический день", "day", "День без уроков: Пн, Вт, Ср, Чт или Пт. Пусто — нет.", "Ср",
               False, WEEKDAYS),
        Column("свой кабинет", "text", "Номер кабинета с листа «Кабинеты», если у учителя свой. Пусто — нет.", "21"),
        Column("уроков в день", "number", "Больше стольких уроков в один день не ставить. Пусто — без ограничения.", ""),
        Column("совместитель", "bool", "да — работает ещё в другой школе.", "нет"),
    ]),
    SheetSpec("Нагрузка", "load", "Кто что ведёт: строка на каждый предмет в каждом классе. При делении "
              "на подгруппы — отдельная строка на каждую подгруппу со своим учителем.", [
        Column("класс", "text", "Как на листе «Классы».", "5А", True),
        Column("предмет", "text", "Как на листе «Предметы».", "Математика", True),
        Column("учитель", "text", "ФИО как на листе «Учителя».", "Иванова И. И.", True),
        Column("часов", "number", "Уроков в неделю.", "5", True),
        Column("подгруппа", "text", "Пусто — весь класс. При делении — 1 и 2, две строки с разными учителями.", ""),
        Column("уровень", "choice", "Уровень изучения. Пусто — базовый.", "базовый", False, list(LEVELS)),
        Column("тип", "choice", "Пусто — обычный урок.", "урок", False, list(LESSON_KINDS)),
        Column("кабинет", "choice", "Только если подгруппы идут в разные кабинеты (труд). Обычно пусто.", "",
               False, ROOMS),
    ]),
]
BY_TABLE = {spec.table: spec for spec in SPEC}
GUIDE_SHEET = "Как заполнять"

RULES = [
    "Заполняйте листы по порядку: Классы, Кабинеты, Предметы, Учителя, Нагрузка.",
    "Названия классов, предметов и ФИО пишите одинаково на всех листах: «5А» на листе «Нагрузка» — "
    "это «5А» с листа «Классы».",
    "Названия листов и колонок не меняйте. Лишние колонки можно оставить — система их пропустит.",
    "Синие заголовки — обязательные колонки. У каждого заголовка есть подсказка: наведите на него.",
    "В колонках «да/нет» пишите «да», «нет», «+» или «-». Пусто — это «нет».",
    "Загружаются только листы, которые есть в файле: остальные данные школы не меняются, а прежние "
    "остаются в истории.",
]


def guide() -> dict:
    """Описание файла для панели «Как оформить Excel» на экране данных."""
    return {"rules": RULES,
            "sheets": [{"sheet": s.sheet, "about": s.about, "columns": [asdict(c) for c in s.columns]}
                       for s in SPEC]}


# ---------------------------------------------------------------- выгрузка

def _blank(value) -> bool:
    return value is None or (isinstance(value, float) and math.isnan(value)) \
        or (isinstance(value, str) and value.strip() in ("", NONE_CHOICE, "nan", "None"))


def _export(column: Column, value):
    if _blank(value):
        return ""
    if column.kind == "bool":
        return "да" if _parse_bool(value)[0] else "нет"
    if column.kind == "day":
        text = str(value).strip()
        return DAY_NAMES.get(int(float(text)), "") if text.replace(".", "", 1).isdigit() else text
    return value


def write_book(tables: dict[str, pd.DataFrame], blank: bool = False) -> bytes:
    """Файл данных школы: лист «Как заполнять» и листы таблиц (пустые, если blank)."""
    book = Workbook()
    _write_guide(book.active)
    for spec in SPEC:
        sheet = book.create_sheet(spec.sheet)
        for c, column in enumerate(spec.columns, start=1):
            letter = get_column_letter(c)
            head = sheet.cell(row=1, column=c, value=column.name)
            head.font = Font(bold=True, color="FFFFFF" if column.required else "1C2B4F")
            head.fill = PatternFill("solid", fgColor="2346B0" if column.required else "E6ECFA")
            note = column.hint + (f"\nНапример: {column.example}" if column.example else "") \
                + ("\nОбязательная колонка." if column.required else "")
            head.comment = Comment(note, "ЛАД", width=260, height=110)
            sheet.column_dimensions[letter].width = max(14, len(column.name) + 6)
            options = ["да", "нет"] if column.kind == "bool" else column.options
            if options:
                # Предупреждение, а не запрет: завуч может вставить значение из своей
                # таблицы — разбор при загрузке поймёт «Спортзал» и «среда» сам.
                rule = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                                      showErrorMessage=True, errorStyle="warning",
                                      errorTitle="Значение не из списка",
                                      error="Лучше выбрать из списка: " + ", ".join(options))
                sheet.add_data_validation(rule)
                rule.add(f"{letter}2:{letter}2000")
        if not blank:
            table = tables.get(spec.table, pd.DataFrame())
            for r, (_, row) in enumerate(table.iterrows(), start=2):
                for c, column in enumerate(spec.columns, start=1):
                    sheet.cell(row=r, column=c, value=_export(column, row.get(column.name, "")))
        sheet.freeze_panes = "A2"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _write_guide(sheet) -> None:
    sheet.title = GUIDE_SHEET
    wrap = Alignment(wrap_text=True, vertical="top")
    sheet["A1"] = "Как заполнять файл данных школы для ЛАД"
    sheet["A1"].font = Font(bold=True, size=16, color="1C2B4F")
    row = 3
    for n, rule in enumerate(RULES, start=1):
        sheet.cell(row=row, column=1, value=f"{n}. {rule}").alignment = Alignment(vertical="top")
        row += 1
    for spec in SPEC:
        row += 1
        sheet.cell(row=row, column=1, value=f"Лист «{spec.sheet}»").font = Font(bold=True, size=13, color="1C2B4F")
        row += 1
        sheet.cell(row=row, column=1, value=spec.about).font = Font(color="5F6A84")
        row += 1
        for c, title in enumerate(["Колонка", "Обязательно", "Что писать", "Пример", "Допустимые значения"], start=1):
            cell = sheet.cell(row=row, column=c, value=title)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E6ECFA")
        for column in spec.columns:
            row += 1
            options = "да, нет" if column.kind == "bool" else ", ".join(column.options)
            for c, value in enumerate([column.name, "да" if column.required else "", column.hint,
                                       column.example, options], start=1):
                sheet.cell(row=row, column=c, value=value).alignment = wrap
        row += 1
    for letter, width in zip("ABCDE", (22, 13, 70, 16, 44)):
        sheet.column_dimensions[letter].width = width


# ---------------------------------------------------------------- загрузка

YES = {"да", "д", "yes", "y", "true", "истина", "1", "+", "x", "х", "v", "✓", "✔"}
NO = {"нет", "н", "no", "n", "false", "ложь", "0", "-", "–", "—", ""}
DAY_BY_TEXT = {"пн": 1, "пон": 1, "понедельник": 1, "вт": 2, "вто": 2, "вторник": 2, "ср": 3, "сре": 3,
               "среда": 3, "чт": 4, "че": 4, "чет": 4, "четверг": 4, "пт": 5, "пя": 5, "пят": 5,
               "пятница": 5, "сб": 6, "суб": 6, "суббота": 6}


def _text(raw) -> str:
    if raw is None or (isinstance(raw, float) and math.isnan(raw)):
        return ""
    if isinstance(raw, float) and raw.is_integer():
        return str(int(raw))  # «1.0» из числовой ячейки — это подгруппа «1»
    return str(raw).strip()


def _parse_bool(raw) -> tuple[bool, str | None]:
    if isinstance(raw, bool):
        return raw, None
    text = _text(raw).lower()
    if text in YES:
        return True, None
    if text in NO:
        return False, None
    return False, f"не понял «{_text(raw)}» — пишите «да» или «нет»"


def _parse(column: Column, raw, default):
    """Значение ячейки → значение таблицы школы и, если не понято, объяснение."""
    text = _text(raw)
    if column.kind == "bool":
        # Пустое «только в нём» — «решит система по типу кабинета», а не «нет».
        if text == "" and column.name == "только в нём":
            return "", None
        return _parse_bool(raw)
    if text in ("", NONE_CHOICE):
        return default if default == NONE_CHOICE else "", None
    if column.kind == "number":
        if isinstance(raw, (int, float)) and not isinstance(raw, bool) and float(raw).is_integer() and raw >= 0:
            return int(raw), None
        try:
            number = float(text.replace(",", ".").replace(" ", ""))
            if number.is_integer() and number >= 0:
                return int(number), None
        except ValueError:
            pass
        return "", f"«{text}» — здесь нужно целое число"
    if column.kind == "day":
        lowered = text.lower().rstrip(".")
        if lowered.replace(".0", "").isdigit() and 1 <= int(float(lowered)) <= 6:
            return str(int(float(lowered))), None
        day = DAY_BY_TEXT.get(lowered) or DAY_BY_TEXT.get(lowered[:3]) or DAY_BY_TEXT.get(lowered[:2])
        if day:
            return str(day), None
        return "", f"не понял день «{text}» — пишите Пн, Вт, Ср, Чт или Пт"
    if column.kind == "choice":
        exact = {o.lower(): o for o in column.options}
        loose = {o.lower().replace(" ", ""): o for o in column.options}
        found = exact.get(text.lower()) or loose.get(text.lower().replace(" ", ""))
        if found:
            return found, None
        return text, f"«{text}» нет в списке: {', '.join(column.options)}"
    return text, None


def normalize(table: str, frame: pd.DataFrame, sheet: str) -> tuple[pd.DataFrame, list[dict]]:
    """Привести значения листа к виду таблиц школы. Номер строки — как в Excel."""
    spec = BY_TABLE[table]
    defaults = blank_tables()[table].iloc[0]
    frame = frame.copy()
    issues: list[dict] = []
    for column in spec.columns:
        if column.name not in frame.columns:
            continue
        values = []
        for index, raw in frame[column.name].items():
            value, problem = _parse(column, raw, defaults.get(column.name))
            if problem:
                issues.append({"sheet": sheet, "row": int(index) + 2, "column": column.name,
                               "value": _text(raw), "message": problem})
            values.append(value)
        frame[column.name] = pd.Series(values, index=frame.index, dtype=object)
    for index, row in frame.iterrows():
        for column in spec.columns:
            if column.required and column.name in frame.columns and _text(row[column.name]) in ("", NONE_CHOICE):
                if not any(i["row"] == int(index) + 2 and i["column"] == column.name for i in issues):
                    issues.append({"sheet": sheet, "row": int(index) + 2, "column": column.name, "value": "",
                                   "message": "пусто, а колонка обязательная"})
    return frame, issues
