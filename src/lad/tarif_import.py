"""Импорт чужой таблицы нагрузки (тарификации) — в каком виде она есть у школы.

Ввод данных — 80% работы завуча (CLAUDE.md §3.3), а нагрузка у школы уже лежит
в Excel: тарификационный список, выгрузка из бухгалтерии, своя табличка. Под наш
шаблон её никто перепечатывать не станет. Поэтому здесь не «загрузите по форме»,
а «покажите, какая колонка что значит» — и система разбирает сама.

Встречаются две раскладки:
  • «строка на урок» — колонки Класс, Предмет, Учитель, Часов;
  • «учитель × классы» — строка на учителя и предмет, классы колонками,
    в ячейке часы (так выглядит большинство тарификационных списков).

Что разбирается без вопросов, но попадает в отчёт:
  • объединённые ячейки: ФИО на несколько строк — во все строки;
  • класс «5 «А»», «5-а», «5A» латиницей → «5А»;
  • ФИО полностью → «Фамилия И. О.», как в расписании;
  • «Матем.», «Рус. яз.», «Физ-ра» → название из учебного плана;
  • строки «Итого» и «Всего» пропускаются;
  • один предмет в классе у двух учителей без подгруппы → подгруппы 1 и 2
    (иначе язык с делением стал бы двойной нагрузкой класса).
Что не понято — не угадывается, а называется со строкой Excel.
"""

from __future__ import annotations

import io
import re
from collections import defaultdict

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .tables import NONE_CHOICE, ROOM_KINDS, STRICT_ROOM_KINDS, load_plan

FIELDS = ("class", "subject", "teacher", "hours", "subgroup")
SYNONYMS = {
    "class": ["класс", "классы", "кл"],
    "subject": ["предмет", "учебный предмет", "наименование предмета", "дисциплина", "предметы"],
    "teacher": ["учитель", "фио", "ф и о", "фио учителя", "преподаватель", "педагог", "учителя"],
    "hours": ["часов в неделю", "часов", "часы", "кол во часов", "количество часов", "нагрузка", "ч нед", "всего часов"],
    "subgroup": ["подгруппа", "подгр", "группа"],
}
LATIN = str.maketrans("ABEKMHOPCTXaekmhopctxb", "АВЕКМНОРСТХаекмнорстхв")
CLASS_RE = re.compile(r"^\s*(\d{1,2})\s*[-–—\s\"«»“”„']*\s*([А-ЯЁа-яёA-Za-z])?\s*[\"»”']*\s*(?:класс|кл\.?)?\s*$")
TOTAL_RE = re.compile(r"(^|\s)(итого|всего)(\s|:|$)", re.I)
ALIASES = {
    "физкультура": "Физическая культура и здоровье", "физ ра": "Физическая культура и здоровье",
    "физра": "Физическая культура и здоровье", "физическая культура": "Физическая культура и здоровье",
    "изо": "Изобразительное искусство", "дмп": "Допризывная и медицинская подготовка",
    "чим": "Человек и мир", "английский язык": "Иностранный язык", "немецкий язык": "Иностранный язык",
    "французский язык": "Иностранный язык", "англ яз": "Иностранный язык", "нем яз": "Иностранный язык",
    "труд": "Трудовое обучение", "трудовое": "Трудовое обучение", "обж": "Основы безопасности жизнедеятельности",
    "мхк": "Искусство (отечественная и мировая художественная культура)",
}


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value)).strip()


def _key(text: str) -> str:
    return " ".join(re.findall(r"[а-яa-z0-9]+", text.lower().replace("ё", "е")))


def class_name(text: str, need_letter: bool = False) -> str | None:
    """«5 «А»», «5-а», «5A» → «5А». need_letter — для заголовков колонок: «5» там это не класс."""
    match = CLASS_RE.match(text or "")
    if not match or not 1 <= int(match.group(1)) <= 11:
        return None
    letter = (match.group(2) or "").translate(LATIN).upper()
    if need_letter and not letter:
        return None
    return f"{int(match.group(1))}{letter}"


def short_name(name: str) -> str:
    """«Иванова Ирина Ивановна» → «Иванова И. И.»; «Иванова И.И.» → «Иванова И. И.»."""
    words = name.replace(".", ". ").split()
    if len(words) == 3 and all(w[:1].isupper() for w in words):
        return f"{words[0]} {words[1][0]}. {words[2][0]}."
    return re.sub(r"\.\s*(?=\S)", ". ", name).strip()


def parse_hours(text: str) -> tuple[int | None, str | None]:
    if not text:
        return None, None
    if "/" in text:
        return None, f"«{text}» — часы по полугодиям, такие вводятся вручную"
    try:
        number = float(text.replace(",", ".").replace(" ", ""))
    except ValueError:
        return None, f"«{text}» — не число часов"
    if not number.is_integer() or number < 0:
        return None, f"«{text}» — часы должны быть целым числом (часы по полугодиям вводятся вручную)"
    return int(number), None


def match_subject(name: str, known: list[str]) -> str | None:
    """Название из учебного плана или из данных школы по сокращению; None — не узнали."""
    key = _key(name)
    for candidate in known:
        if _key(candidate) == key:
            return candidate
    if key in ALIASES:
        return ALIASES[key]
    tokens = key.split()
    if not tokens:
        return None
    found = [c for c in known if len(_key(c).split()) >= len(tokens)
             and all(word.startswith(token) for token, word in zip(tokens, _key(c).split()))]
    return found[0] if len(found) == 1 else None


# ---------------------------------------------------------------- чтение и догадка

def read_book(data: bytes) -> dict[str, list[list[str]]]:
    try:
        book = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as error:  # noqa: BLE001
        raise ValueError("Файл не читается как Excel (.xlsx). Если это старый .xls — откройте его в Excel "
                         "и сохраните как «Книга Excel (.xlsx)».") from error
    sheets = {}
    for sheet in book.worksheets:
        merged = {}
        for area in sheet.merged_cells.ranges:
            value = sheet.cell(area.min_row, area.min_col).value
            # Только вниз по первой колонке: ФИО на несколько строк нужно в каждой строке,
            # а заголовок на всю ширину («Тарификационный список…») не должен
            # размножаться по колонкам классов.
            for r in range(area.min_row, area.max_row + 1):
                merged[r, area.min_col] = value
        rows = []
        for r, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            rows.append([_cell(merged.get((r, c), value)) for c, value in enumerate(row, start=1)])
            if r >= 3000:
                break
        while rows and not any(rows[-1]):
            rows.pop()
        width = max((max((i + 1 for i, v in enumerate(row) if v), default=0) for row in rows), default=0)
        sheets[sheet.title] = [(row + [""] * width)[:width] for row in rows]
    return sheets


def guess(rows: list[list[str]]) -> dict:
    """Строка заголовков, раскладка и колонки — по названиям в первых 20 строках."""
    best = {"score": 0, "header_row": 0, "layout": "long", "mapping": {f: None for f in FIELDS}, "class_columns": []}
    for index, row in enumerate(rows[:20]):
        keys = [_key(text) for text in row]
        mapping: dict[str, int | None] = {f: None for f in FIELDS}
        taken: set[int] = set()
        for field in FIELDS:
            for word in SYNONYMS[field]:
                column = next((c for c, k in enumerate(keys) if c not in taken and k and (k == word or k.startswith(word + " ")
                                                                                         or (len(word) > 3 and k.startswith(word)))), None)
                if column is not None:
                    mapping[field] = column
                    taken.add(column)
                    break
        class_columns = [c for c, text in enumerate(row) if c not in taken and class_name(text, need_letter=True)]
        if len(class_columns) >= 3:
            mapping["class"] = mapping["hours"] = mapping["subgroup"] = None
            score, layout = len(class_columns) + 2 * sum(mapping[f] is not None for f in ("teacher", "subject")), "wide"
        else:
            score, layout = 2 * sum(v is not None for v in mapping.values()), "long"
            class_columns = []
        if score > best["score"]:
            best = {"score": score, "header_row": index, "layout": layout, "mapping": mapping, "class_columns": class_columns}
    best.pop("score")
    return best


def inspect(data: bytes) -> dict:
    sheets = read_book(data)
    return {"sheets": [{"name": name, "rows": rows[:200], "total_rows": len(rows), "guess": guess(rows)}
                       for name, rows in sheets.items() if rows]}


# ---------------------------------------------------------------- разбор

def extract(rows: list[list[str]], header_row: int, layout: str, mapping: dict, class_columns: list[int],
            fill_down: bool = True, shorten: bool = True, known_subjects: list[str] | None = None) -> dict:
    """Строки нагрузки из таблицы школы и отчёт: что узнано, что пропущено и почему."""
    known = list(dict.fromkeys((known_subjects or []) + [s["name"] for s in load_plan().get("subjects", [])]))
    header = rows[header_row] if header_row < len(rows) else []
    get = lambda row, field: row[mapping[field]] if mapping.get(field) is not None and mapping[field] < len(row) else ""  # noqa: E731
    records, skipped, matched, unknown_subjects = [], [], {}, set()
    last_teacher = last_subject = ""

    for r in range(header_row + 1, len(rows)):
        row, excel_row = rows[r], r + 1
        if not any(row):
            continue
        if any(TOTAL_RE.search(text) for text in row if text):
            skipped.append({"row": excel_row, "reason": "итоговая строка"})
            continue
        teacher, subject = get(row, "teacher"), get(row, "subject")
        if fill_down:
            if teacher and teacher != last_teacher:
                last_subject = ""
            teacher = teacher or last_teacher
            subject = subject or last_subject
            last_teacher, last_subject = teacher, subject
        if not subject:
            skipped.append({"row": excel_row, "reason": "нет предмета"})
            continue
        name = match_subject(subject, known)
        if name and name != subject:
            matched[subject] = name
        if not name:
            unknown_subjects.add(subject)
        subject_name = name or subject
        teacher_name = short_name(teacher) if shorten and teacher else teacher

        if layout == "wide":
            for column in class_columns:
                value = row[column] if column < len(row) else ""
                hours, problem = parse_hours(value)
                if problem:
                    skipped.append({"row": excel_row, "reason": f"колонка {header[column]}: {problem}"})
                if hours:
                    records.append({"класс": class_name(header[column], need_letter=True), "предмет": subject_name,
                                    "учитель": teacher_name, "часов": hours, "подгруппа": "", "row": excel_row})
            continue

        raw_class = get(row, "class")
        klass = class_name(raw_class)
        hours, problem = parse_hours(get(row, "hours"))
        if not klass:
            skipped.append({"row": excel_row, "reason": f"не понял класс «{raw_class}»" if raw_class else "нет класса"})
            continue
        if problem or hours is None:
            skipped.append({"row": excel_row, "reason": problem or "нет часов"})
            continue
        if hours == 0:
            continue
        records.append({"класс": klass, "предмет": subject_name, "учитель": teacher_name, "часов": hours,
                        "подгруппа": get(row, "subgroup"), "row": excel_row})

    # Один предмет в классе у двух учителей без подгрупп — это деление класса.
    split = []
    by_lesson = defaultdict(list)
    for record in records:
        if not record["подгруппа"]:
            by_lesson[record["класс"], record["предмет"]].append(record)
    for (klass, subject), items in by_lesson.items():
        teachers = list(dict.fromkeys(item["учитель"] for item in items))
        if len(items) > 1 and len(teachers) == len(items):
            for n, item in enumerate(items, start=1):
                item["подгруппа"] = str(n)
            split.append(f"{klass} «{subject}»: {', '.join(teachers)}")

    return {"records": records, "skipped": skipped, "matched": [{"from": k, "to": v} for k, v in matched.items()],
            "unknown_subjects": sorted(unknown_subjects), "split": split,
            "no_teacher": sum(1 for r in records if not r["учитель"])}


def apply(tables: dict[str, pd.DataFrame], records: list[dict], mode: str) -> tuple[dict[str, pd.DataFrame], dict]:
    """Положить нагрузку в таблицы школы; недостающие классы, предметы и учителей завести."""
    tables = {name: table.copy() for name, table in tables.items()}
    plan_rooms = {s["name"]: s.get("room", "обычный") for s in load_plan().get("subjects", [])}

    def names(table: str, column: str) -> set[str]:
        frame = tables[table]
        return set(frame[column].astype(str).str.strip()) if column in frame else set()

    def add(table: str, rows: list[dict]) -> None:
        if rows:
            tables[table] = pd.concat([tables[table], pd.DataFrame(rows)], ignore_index=True)

    new_classes = sorted({r["класс"] for r in records} - names("classes", "класс"),
                         key=lambda n: (int(re.match(r"\d+", n).group()), n))
    new_subjects = sorted({r["предмет"] for r in records} - names("subjects", "предмет"))
    new_teachers = sorted({r["учитель"] for r in records if r["учитель"]} - names("teachers", "ФИО"))
    add("classes", [{"класс": n, "учеников": 24, "повышенный уровень": False} for n in new_classes])
    add("subjects", [{"предмет": n, "кабинет": plan_rooms.get(n, "обычный"),
                      "только в нём": ROOM_KINDS.get(plan_rooms.get(n, "обычный")) in STRICT_ROOM_KINDS,
                      "всегда парой": False} for n in new_subjects])
    add("teachers", [{"ФИО": n, "методический день": "", "свой кабинет": NONE_CHOICE, "уроков в день": "",
                      "совместитель": False} for n in new_teachers])

    load = [{"класс": r["класс"], "предмет": r["предмет"], "учитель": r["учитель"], "часов": r["часов"],
             "подгруппа": r["подгруппа"], "уровень": "базовый", "тип": "урок", "кабинет": NONE_CHOICE} for r in records]
    frame = pd.DataFrame(load)
    if mode == "add" and len(tables["load"]):
        frame = pd.concat([tables["load"], frame], ignore_index=True)
        # та же строка (класс, предмет, учитель, подгруппа) — берём часы из файла
        frame = frame.drop_duplicates(subset=["класс", "предмет", "учитель", "подгруппа"], keep="last")
    tables["load"] = frame.reset_index(drop=True)
    return tables, {"classes": new_classes, "subjects": new_subjects, "teachers": new_teachers,
                    "load_rows": len(tables["load"]), "mode": mode}


def column_label(index: int, header: list[str]) -> str:
    title = header[index] if index < len(header) else ""
    return f"{get_column_letter(index + 1)}: {title}" if title else get_column_letter(index + 1)
