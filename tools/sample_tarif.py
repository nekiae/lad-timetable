"""Пример чужой таблицы тарификации — для кадра «импорт своей таблицы» в саммари.

    .venv/bin/python tools/sample_tarif.py

Кладёт data/summary/tarif-example.xlsx. Таблица нарочно оформлена так, как её
делают в школе: шапка не в первой строке, ФИО объединено на несколько строк,
предметы сокращены, в конце «Итого». Имена вымышленные — лист уходит наружу.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "summary" / "tarif-example.xlsx"

ROWS = [
    ("Иванова А. С.", [("Матем.", "5А", 5), ("Матем.", "5Б", 5), ("Матем.", "6А", 5)]),
    ("Петрова Н. В.", [("Рус. яз.", "5А", 4), ("Рус. лит.", "5А", 2), ("Рус. яз.", "5Б", 4)]),
    ("Сидоров К. П.", [("Физ-ра", "5А", 3), ("Физ-ра", "5Б", 3), ("Физ-ра", "6А", 3)]),
]


def build() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Тарификация"

    ws["A1"] = "Тарификация педагогических работников на 2026/2027 учебный год"
    ws["A1"].font = Font(bold=True, size=12)
    ws.merge_cells("A1:D1")

    head = ["Ф.И.О. учителя", "Предмет", "Класс", "Часов в неделю"]
    for col, name in enumerate(head, start=1):
        cell = ws.cell(row=3, column=col, value=name)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    row = 4
    total = 0
    for teacher, lessons in ROWS:
        first = row
        for subject, klass, hours in lessons:
            ws.cell(row=row, column=1, value=teacher if row == first else None)
            ws.cell(row=row, column=2, value=subject)
            ws.cell(row=row, column=3, value=klass)
            ws.cell(row=row, column=4, value=hours)
            total += hours
            row += 1
        # ФИО объединено по строкам учителя — обычное оформление тарификации.
        ws.merge_cells(start_row=first, start_column=1, end_row=row - 1, end_column=1)
        ws.cell(row=first, column=1).alignment = Alignment(vertical="center")

    ws.cell(row=row, column=1, value="Итого").font = Font(bold=True)
    ws.cell(row=row, column=4, value=total).font = Font(bold=True)

    for col, width in zip("ABCD", (24, 14, 10, 16)):
        ws.column_dimensions[col].width = width

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    print("готово:", build())
