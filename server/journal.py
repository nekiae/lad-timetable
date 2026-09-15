"""Журнал замен: сохранённые замены по датам и итог за месяц.

Лист замен раньше жил до перезагрузки страницы. А школе нужна история: кто кого
и когда заменял — по ней в конце месяца считают часы замен к оплате, и по ней
же видно, кого завуч нагружает заменами чаще других.

Одна запись — один отсутствующий учитель в одну дату. Повторное сохранение
той же даты и учителя обновляет запись, а не плодит дубли.
"""

from __future__ import annotations

import re
from collections import Counter
from io import BytesIO

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from pydantic import BaseModel

from . import db

router = APIRouter(prefix="/api/schools/{school_id}/substitutions/journal")

NOT_HELD = "урок не проводится"
WEEKDAYS = ["понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье"]


class JournalRow(BaseModel):
    period: int
    group: str
    subject: str
    room: str = ""
    substitute: str = ""  # ФИО заменяющего; пусто — урок не проводится


class JournalEntry(BaseModel):
    date: str  # YYYY-MM-DD
    absent: str  # ФИО отсутствующего
    rows: list[JournalRow]


def _school(school_id: str) -> None:
    if not db.get_school(school_id):
        raise HTTPException(404, "Школа не найдена")


def _month(month: str) -> str:
    if not re.fullmatch(r"\d{4}-\d{2}", month or ""):
        raise HTTPException(422, "Месяц в виде ГГГГ-ММ")
    return month


def summary(entries: list[dict]) -> dict:
    """Итог месяца: сколько уроков провёл каждый заменяющий и пропустил каждый отсутствующий."""
    by_substitute: Counter = Counter()
    by_absent: Counter = Counter()
    not_held = 0
    for entry in entries:
        for row in entry["rows"]:
            by_absent[entry["absent"]] += 1
            if row.get("substitute"):
                by_substitute[row["substitute"]] += 1
            else:
                not_held += 1
    order = lambda c: sorted(({"name": n, "lessons": k} for n, k in c.items()),  # noqa: E731
                             key=lambda x: (-x["lessons"], x["name"]))
    return {"by_substitute": order(by_substitute), "by_absent": order(by_absent),
            "total": sum(by_substitute.values()), "not_held": not_held}


@router.get("")
def month_journal(school_id: str, month: str) -> dict:
    _school(school_id)
    entries = db.list_substitutions(school_id, _month(month))
    return {"month": month, "days": entries, **summary(entries)}


@router.put("")
def save_entry(school_id: str, body: JournalEntry) -> dict:
    _school(school_id)
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", body.date) or not body.absent.strip():
        raise HTTPException(422, "Нужны дата и отсутствующий учитель")
    return db.save_substitution(school_id, body.date, body.absent.strip(), [r.model_dump() for r in body.rows])


@router.delete("/{entry_id}")
def delete_entry(school_id: str, entry_id: str) -> dict:
    _school(school_id)
    if not db.delete_substitution(school_id, entry_id):
        raise HTTPException(404, "Запись не найдена")
    return {"ok": True}


@router.get(".xlsx")
def month_xlsx(school_id: str, month: str) -> Response:
    """Журнал за месяц в Excel: все замены построчно и итог по учителям — для учёта часов."""
    _school(school_id)
    entries = db.list_substitutions(school_id, _month(month))
    total = summary(entries)
    book = Workbook()
    line = Side(style="thin", color="BFBFBF")

    def table(sheet, title: str, head: list[str], rows: list[list], widths: list[int]) -> None:
        sheet["A1"] = title
        sheet["A1"].font = Font(bold=True, size=14)
        for c, name in enumerate(head, start=1):
            cell = sheet.cell(row=3, column=c, value=name)
            cell.font = Font(bold=True)
            cell.border = Border(bottom=Side(style="medium", color="1C2B4F"))
        for r, values in enumerate(rows, start=4):
            for c, value in enumerate(values, start=1):
                cell = sheet.cell(row=r, column=c, value=value)
                cell.border = Border(bottom=line)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for c, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + c)].width = width
        sheet.freeze_panes = "A4"

    from datetime import date as _date
    lines = []
    for entry in entries:
        day = _date.fromisoformat(entry["date"])
        for row in sorted(entry["rows"], key=lambda x: x["period"]):
            lines.append([day.strftime("%d.%m.%Y"), WEEKDAYS[day.weekday()], row["period"], row["group"],
                          row["subject"], row.get("room", ""), entry["absent"], row.get("substitute") or NOT_HELD])
    sheet = book.active
    sheet.title = "Замены"
    table(sheet, f"Журнал замен за {month}", ["Дата", "День", "Урок", "Класс", "Предмет", "Кабинет",
                                             "Отсутствовал", "Заменял"], lines, [12, 13, 7, 9, 30, 9, 24, 24])
    table(book.create_sheet("Кто заменял"), f"Проведено замен за {month}", ["Учитель", "Уроков"],
          [[x["name"], x["lessons"]] for x in total["by_substitute"]]
          + [["Всего", total["total"]], ["Не проводилось", total["not_held"]]], [30, 10])
    table(book.create_sheet("Кто отсутствовал"), f"Пропущено уроков за {month}", ["Учитель", "Уроков"],
          [[x["name"], x["lessons"]] for x in total["by_absent"]], [30, 10])
    out = BytesIO()
    book.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="zameny-{month}.xlsx"'})
